"""Generate and validate an RTL DOCX from the final Persian reviewer responses."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = PROJECT_ROOT / ".docs" / "07_responses_to_reviewers_fa.md"
DEFAULT_OUTPUT = PROJECT_ROOT / ".docs" / "07_responses_to_reviewers_fa.docx"
FONT_NAME = "Tahoma"


def source_blocks(markdown: str) -> list[tuple[int, str]]:
    blocks: list[tuple[int, str]] = []
    for raw in re.split(r"\n\s*\n", markdown.strip()):
        text = " ".join(line.strip() for line in raw.splitlines()).strip()
        if text.startswith("# "):
            blocks.append((1, text[2:]))
        elif text.startswith("## "):
            blocks.append((2, text[3:]))
        else:
            blocks.append((0, text.replace("**", "")))
    return blocks


def set_run_font(run, size: float, *, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)
    rtl = OxmlElement("w:rtl")
    run_properties.append(rtl)


def set_rtl(paragraph, *, space_before: float = 0, space_after: float = 6) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.25
    properties = paragraph._p.get_or_add_pPr()
    bidi = properties.find(qn("w:bidi"))
    if bidi is None:
        properties.append(OxmlElement("w:bidi"))


def add_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    set_rtl(paragraph, space_after=0)
    prefix = paragraph.add_run("صفحه ")
    set_run_font(prefix, 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.extend([begin, instruction, separate, value, end])
    set_run_font(field_run, 9)


def build(source: Path, output: Path) -> None:
    blocks = source_blocks(source.read_text(encoding="utf-8"))
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    add_footer(section)

    document.core_properties.title = blocks[0][1]
    document.core_properties.subject = "پاسخ مستند به نظرات داوران پروژه طبقه‌بندی برانگیختگی DREAMER"
    document.core_properties.author = "StressClassification Research Project"

    for level, text in blocks:
        paragraph = document.add_paragraph()
        if level == 1:
            set_rtl(paragraph, space_before=0, space_after=18)
            run = paragraph.add_run(text)
            set_run_font(run, 18, bold=True)
        elif level == 2:
            set_rtl(paragraph, space_before=14, space_after=8)
            run = paragraph.add_run(text)
            set_run_font(run, 14, bold=True)
            properties = paragraph._p.get_or_add_pPr()
            keep_next = OxmlElement("w:keepNext")
            properties.append(keep_next)
        else:
            set_rtl(paragraph, space_before=0, space_after=7)
            label_match = re.match(r"^(ابهام:|پاسخ:|اصلاح انجام شده:|محل اصلاح:)(.*)$", text)
            if label_match:
                label = paragraph.add_run(label_match.group(1))
                set_run_font(label, 11, bold=True)
                body = paragraph.add_run(label_match.group(2))
                set_run_font(body, 11)
            else:
                run = paragraph.add_run(text)
                set_run_font(run, 11)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    validate(source, output)


def validate(source: Path, output: Path) -> None:
    expected = [text for _, text in source_blocks(source.read_text(encoding="utf-8"))]
    document = Document(output)
    actual = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    if actual != expected:
        raise RuntimeError("DOCX text differs from Markdown or is reordered")
    for paragraph in document.paragraphs:
        if paragraph.text and paragraph._p.pPr.find(qn("w:bidi")) is None:
            raise RuntimeError(f"paragraph is not RTL: {paragraph.text[:40]}")
        if paragraph.text and paragraph.alignment != WD_ALIGN_PARAGRAPH.RIGHT:
            raise RuntimeError(f"paragraph is not right-aligned: {paragraph.text[:40]}")
    if output.stat().st_size < 20_000:
        raise RuntimeError("DOCX is unexpectedly small")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
    print(f"generated and validated {args.output}")


if __name__ == "__main__":
    main()
