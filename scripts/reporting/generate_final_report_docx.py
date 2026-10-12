"""Build and validate the Persian B Nazanin final experimental report."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from persian_docx import (
    FONT_NAME,
    add_mixed_text,
    add_page_number,
    add_toc,
    plain_markdown,
    prevent_row_split,
    repeat_table_header,
    set_cell_rtl,
    set_paragraph_rtl,
    set_table_rtl,
    shade_cell,
    validate_b_nazanin,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = PROJECT_ROOT / ".docs" / "09_final_experimental_report_fa.md"
DEFAULT_OUTPUT = PROJECT_ROOT / ".docs" / "09_final_experimental_report_fa.docx"
RESULTS = PROJECT_ROOT / "results" / "corrected" / "dreamer_arousal_3class"
IMAGE_RE = re.compile(r"^!\[(.+?)\]\((.+?)\)$")
TABLE_SEPARATOR_RE = re.compile(r"^:?-{3,}:?$")
FA_TRANSLATION = str.maketrans("0123456789.", "۰۱۲۳۴۵۶۷۸۹٫")


def fa_number(value: float, digits: int = 4) -> str:
    return f"{value:.{digits}f}".translate(FA_TRANSLATION)


def parse_markdown(source: Path) -> list[dict[str, object]]:
    lines = source.read_text(encoding="utf-8").splitlines()
    blocks: list[dict[str, object]] = []
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("### "):
            blocks.append({"type": "heading", "level": 3, "text": plain_markdown(stripped[4:])})
            index += 1
            continue
        if stripped.startswith("## "):
            blocks.append({"type": "heading", "level": 2, "text": plain_markdown(stripped[3:])})
            index += 1
            continue
        if stripped.startswith("# "):
            blocks.append({"type": "heading", "level": 1, "text": plain_markdown(stripped[2:])})
            index += 1
            continue
        image = IMAGE_RE.match(stripped)
        if image:
            blocks.append({"type": "image", "alt": image.group(1), "path": image.group(2)})
            index += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [plain_markdown(cell.strip()) for cell in lines[index].strip().strip("|").split("|")]
                if not all(TABLE_SEPARATOR_RE.fullmatch(cell) for cell in cells):
                    rows.append(cells)
                index += 1
            blocks.append({"type": "table", "rows": rows})
            continue
        if stripped.startswith("- "):
            while index < len(lines) and lines[index].strip().startswith("- "):
                blocks.append({"type": "bullet", "text": plain_markdown(lines[index].strip()[2:])})
                index += 1
            continue
        if stripped.startswith("*") and stripped.endswith("*"):
            blocks.append({"type": "caption", "text": plain_markdown(stripped)})
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and lines[index].strip():
            candidate = lines[index].strip()
            if (
                candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith("- ")
                or IMAGE_RE.match(candidate)
                or (candidate.startswith("*") and candidate.endswith("*"))
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        blocks.append({"type": "paragraph", "text": plain_markdown(" ".join(paragraph_lines))})
    return blocks


def configure_styles(document: Document) -> None:
    definitions = {
        "Normal": (12, False),
        "Title": (22, True),
        "Heading 1": (16, True),
        "Heading 2": (14, True),
        "Heading 3": (13, True),
        "Caption": (10, False),
    }
    for name, (size, bold) in definitions.items():
        style = document.styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.bold = bold
        properties = style.element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), FONT_NAME)


def enable_field_updates(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    set_paragraph_rtl(header, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_mixed_text(header, "گزارش نهایی طبقه‌بندی چندوجهی برانگیختگی DREAMER", 9)
    add_page_number(section.footer.paragraphs[0])


def add_cover(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_rtl(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=105,
        space_after=26,
        line_spacing=1.2,
    )
    add_mixed_text(paragraph, title, 23, bold=True)
    subtitle = document.add_paragraph()
    set_paragraph_rtl(subtitle, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
    add_mixed_text(subtitle, "پروژه StressClassification — نسخه نهایی تحویل", 15, bold=True)
    note = document.add_paragraph()
    set_paragraph_rtl(note, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_mixed_text(note, "طبقه‌بندی سه‌کلاسه arousal خوداظهاری با EEG، ECG و HRV", 13)
    note = document.add_paragraph()
    set_paragraph_rtl(note, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_mixed_text(note, "ارزیابی مستقل از آزمودنی — مجموعه‌داده DREAMER", 13)
    document.add_page_break()
    toc_title = document.add_paragraph()
    set_paragraph_rtl(toc_title, alignment=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)
    add_mixed_text(toc_title, "فهرست مطالب", 16, bold=True)
    add_toc(document.add_paragraph())
    document.add_page_break()


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level - 1}" if level > 1 else "Title")
    set_paragraph_rtl(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        space_before=15 if level == 2 else 10,
        space_after=8,
        line_spacing=1.15,
    )
    paragraph.paragraph_format.keep_with_next = True
    add_mixed_text(paragraph, text, 16 if level == 2 else 14, bold=True)


def add_body(document: Document, text: str, *, bullet: bool = False) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_rtl(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=7,
        line_spacing=1.35,
    )
    if bullet:
        paragraph.paragraph_format.right_indent = Cm(0.5)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        add_mixed_text(paragraph, "• " + text, 12)
    else:
        label = re.match(r"^(موضوع:|پروتکل ارزیابی:|معیار اصلی:|محدوده ادعا:)(.*)$", text)
        if label:
            add_mixed_text(paragraph, label.group(1), 12, bold=True)
            add_mixed_text(paragraph, label.group(2), 12)
        else:
            add_mixed_text(paragraph, text, 12)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Caption")
    set_paragraph_rtl(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=2,
        space_after=10,
        line_spacing=1.1,
    )
    paragraph.paragraph_format.keep_with_next = False
    add_mixed_text(paragraph, text, 10, bold=True)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows or len(rows) < 2:
        raise ValueError("report tables require a header and at least one data row")
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise ValueError("inconsistent Markdown table width")
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    set_table_rtl(table)
    font_size = 8.5 if width >= 6 else 9.5
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            repeat_table_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_rtl(cell)
            paragraph = cell.paragraphs[0]
            set_paragraph_rtl(
                paragraph,
                alignment=WD_ALIGN_PARAGRAPH.CENTER if column_index else WD_ALIGN_PARAGRAPH.RIGHT,
                space_after=0,
                line_spacing=1.05,
            )
            add_mixed_text(paragraph, value, font_size, bold=row_index == 0)
            if row_index == 0:
                shade_cell(cell, "D9EAF7")
            elif row_index % 2 == 0:
                shade_cell(cell, "F3F6F8")
    spacer = document.add_paragraph()
    set_paragraph_rtl(spacer, alignment=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)


def add_figure(document: Document, source: Path, alt: str) -> None:
    if not source.is_file():
        raise FileNotFoundError(f"missing report figure: {source}")
    paragraph = document.add_paragraph()
    set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
    run = paragraph.add_run()
    run.add_picture(str(source), width=Cm(15.6))
    inline = run._r.xpath(".//wp:docPr")
    if inline:
        inline[0].set("title", alt)
        inline[0].set("descr", alt)


def validate_measurements(source_text: str) -> None:
    table_rows = []
    for line in source_text.splitlines():
        stripped = line.strip()
        if stripped.startswith("|"):
            cells = [plain_markdown(cell.strip()) for cell in stripped.strip("|").split("|")]
            if not all(TABLE_SEPARATOR_RE.fullmatch(cell) for cell in cells):
                table_rows.append(cells)
    comparison = pd.read_csv(RESULTS / "comparison_all.csv").set_index("model")
    names = {
        "paper_2023_adapted_stacking": "Adapted Stacking",
        "paper_inspired_attention": "Paper-inspired Attention",
        "knn": "kNN",
        "random_forest": "Random Forest",
        "logistic_regression": "Logistic Regression",
        "svm_rbf": "RBF SVM",
        "base_cnn": "Base CNN",
    }
    for model, display in names.items():
        row = comparison.loc[model]
        values = [
            fa_number(row["test_accuracy"]),
            fa_number(row["test_balanced_accuracy"]),
            fa_number(row["test_f1_macro"]),
            fa_number(row["test_roc_auc_ovr_macro"]),
        ]
        matches = [row for row in table_rows if len(row) == 6 and row[0] == display]
        if not any(row[2:] == values for row in matches):
            raise RuntimeError(f"final report does not contain verified comparison values for {model}")

    per_class = pd.read_csv(RESULTS / "per_class_metrics.csv")
    class_names = {"low": "کم", "medium": "متوسط", "high": "زیاد"}
    for _, row in per_class.iterrows():
        expected = [
            names[row["model"]],
            class_names[row["class"]],
            fa_number(row["precision"]),
            fa_number(row["recall"]),
            fa_number(row["f1"]),
            str(int(row["support"])).translate(FA_TRANSLATION),
        ]
        if expected not in table_rows:
            raise RuntimeError(f"missing verified per-class row {row['model']}/{row['class']}")

    model_paths = {
        "knn": RESULTS / "classical/knn/metrics.json",
        "random_forest": RESULTS / "classical/random_forest/metrics.json",
        "logistic_regression": RESULTS / "classical/logistic_regression/metrics.json",
        "svm_rbf": RESULTS / "classical/svm_rbf/metrics.json",
        "paper_2023_adapted_stacking": RESULTS / "paper_2023_adapted_stacking/metrics.json",
        "base_cnn": RESULTS / "base_cnn/metrics.json",
        "paper_inspired_attention": RESULTS / "paper_inspired_attention/metrics.json",
    }
    for model, path in model_paths.items():
        metrics = json.loads(path.read_text(encoding="utf-8"))["trial_level"]
        expected = [
            names[model],
            fa_number(metrics["accuracy"]),
            fa_number(metrics["balanced_accuracy"]),
            fa_number(metrics["f1_macro"]),
            fa_number(metrics["roc_auc_ovr_macro"]),
        ]
        if expected not in table_rows:
            raise RuntimeError(f"missing verified trial-level row for {model}")


def build(source: Path, output: Path) -> None:
    source_text = source.read_text(encoding="utf-8")
    validate_measurements(source_text)
    blocks = parse_markdown(source)
    title = next(block["text"] for block in blocks if block["type"] == "heading" and block["level"] == 1)

    document = Document()
    configure_styles(document)
    enable_field_updates(document)
    section = document.sections[0]
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    add_header_footer(document)
    document.core_properties.title = str(title)
    document.core_properties.subject = "گزارش نهایی آزمایش‌های چندوجهی DREAMER"
    document.core_properties.author = "StressClassification Research Project"
    add_cover(document, str(title))

    for block in blocks:
        kind = block["type"]
        if kind == "heading" and block["level"] == 1:
            continue
        if kind == "heading":
            add_heading(document, str(block["text"]), int(block["level"]))
        elif kind == "paragraph":
            add_body(document, str(block["text"]))
        elif kind == "bullet":
            add_body(document, str(block["text"]), bullet=True)
        elif kind == "caption":
            add_caption(document, str(block["text"]))
        elif kind == "table":
            add_table(document, block["rows"])
        elif kind == "image":
            add_figure(document, (source.parent / str(block["path"])).resolve(), str(block["alt"]))
        else:
            raise ValueError(f"unsupported report block: {kind}")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    validate(source, output)


def validate(source: Path, output: Path) -> None:
    blocks = parse_markdown(source)
    document = Document(output)
    all_text = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        all_text.extend(cell.text for row in table.rows for cell in row.cells if cell.text)

    required = []
    for block in blocks:
        if block["type"] in {"heading", "paragraph", "caption"}:
            required.append(str(block["text"]))
        elif block["type"] == "bullet":
            required.append("• " + str(block["text"]))
        elif block["type"] == "table":
            required.extend(cell for row in block["rows"] for cell in row)
    missing = [text for text in required if text not in all_text]
    if missing:
        raise RuntimeError(f"DOCX is missing or reordered report content: {missing[:5]}")
    if len(document.tables) != 9:
        raise RuntimeError(f"expected 9 scientific tables, found {len(document.tables)}")
    if len(document.inline_shapes) != 7:
        raise RuntimeError(f"expected 7 experimental figures, found {len(document.inline_shapes)}")
    for paragraph in document.paragraphs:
        if (paragraph.text or paragraph._p.xpath(".//w:drawing")) and paragraph._p.pPr.find(qn("w:bidi")) is None:
            raise RuntimeError(f"paragraph is not RTL: {paragraph.text[:40]}")
    for table in document.tables:
        if table._tbl.tblPr.find(qn("w:bidiVisual")) is None:
            raise RuntimeError("a report table is not configured as RTL")
    validate_b_nazanin(document, output)
    if output.stat().st_size < 500_000:
        raise RuntimeError("final report DOCX is unexpectedly small")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
    print(f"generated and validated {args.output}")


if __name__ == "__main__":
    main()
