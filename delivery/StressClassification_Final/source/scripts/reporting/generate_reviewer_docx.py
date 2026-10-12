"""Generate and validate the final B Nazanin RTL reviewer-response DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from persian_docx import (
    FONT_NAME,
    add_mixed_text,
    add_page_number,
    plain_markdown,
    set_paragraph_rtl,
    validate_b_nazanin,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = PROJECT_ROOT / ".docs" / "07_responses_to_reviewers_fa.md"
DEFAULT_OUTPUT = PROJECT_ROOT / ".docs" / "07_responses_to_reviewers_fa.docx"


def source_blocks(markdown: str) -> list[tuple[int, str]]:
    blocks: list[tuple[int, str]] = []
    for raw in re.split(r"\n\s*\n", markdown.strip()):
        text = " ".join(line.strip() for line in raw.splitlines()).strip()
        if text.startswith("# "):
            blocks.append((1, plain_markdown(text[2:])))
        elif text.startswith("## "):
            blocks.append((2, plain_markdown(text[3:])))
        else:
            blocks.append((0, plain_markdown(text)))
    return blocks


def configure_styles(document: Document) -> None:
    for name, size in (("Normal", 12), ("Title", 20), ("Heading 1", 15), ("Heading 2", 13)):
        style = document.styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        properties = style.element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), FONT_NAME)


def build(source: Path, output: Path) -> None:
    blocks = source_blocks(source.read_text(encoding="utf-8"))
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    add_page_number(section.footer.paragraphs[0])

    document.core_properties.title = blocks[0][1]
    document.core_properties.subject = "پاسخ مستند به نظرات داوران پروژه طبقه‌بندی برانگیختگی DREAMER"
    document.core_properties.author = "StressClassification Research Project"

    for level, text in blocks:
        paragraph = document.add_paragraph()
        if level == 1:
            set_paragraph_rtl(
                paragraph,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                space_before=60,
                space_after=24,
                line_spacing=1.15,
            )
            add_mixed_text(paragraph, text, 20, bold=True)
        elif level == 2:
            set_paragraph_rtl(
                paragraph,
                alignment=WD_ALIGN_PARAGRAPH.RIGHT,
                space_before=15,
                space_after=9,
                line_spacing=1.15,
            )
            paragraph.paragraph_format.keep_with_next = True
            add_mixed_text(paragraph, text, 14, bold=True)
        else:
            set_paragraph_rtl(
                paragraph,
                alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                space_before=0,
                space_after=8,
                line_spacing=1.35,
            )
            label_match = re.match(
                r"^(ابهام:|پاسخ:|اصلاح انجام شده:|محل اصلاح:)(.*)$", text
            )
            if label_match:
                add_mixed_text(paragraph, label_match.group(1), 12, bold=True)
                add_mixed_text(paragraph, label_match.group(2), 12)
            else:
                add_mixed_text(paragraph, text, 12)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    validate(source, output)


def validate(source: Path, output: Path) -> None:
    expected = [text for _, text in source_blocks(source.read_text(encoding="utf-8"))]
    document = Document(output)
    actual = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    if actual != expected:
        raise RuntimeError("DOCX text differs from Markdown or is reordered")
    source_text = source.read_text(encoding="utf-8")
    if re.search(r"2023|PLOS|۲۰۲۳", source_text, flags=re.IGNORECASE):
        raise RuntimeError("reviewer-facing response still references the mistaken 2023 paper")
    if "An attention-based multimodal deep learning framework" not in source_text:
        raise RuntimeError("the requested 2026 paper is not identified")
    comment_sections = [text for level, text in source_blocks(source_text) if level == 2 and text.startswith("نظر شماره")]
    if len(comment_sections) != 8:
        raise RuntimeError(f"expected 8 reviewer-comment sections, found {len(comment_sections)}")
    for label in ("ابهام:", "پاسخ:", "اصلاح انجام شده:", "محل اصلاح:"):
        if sum(text.startswith(label) for text in expected) != 8:
            raise RuntimeError(f"expected 8 occurrences of {label}")
    for paragraph in document.paragraphs:
        if paragraph.text and paragraph._p.pPr.find(qn("w:bidi")) is None:
            raise RuntimeError(f"paragraph is not RTL: {paragraph.text[:40]}")
    validate_b_nazanin(document, output)
    if output.stat().st_size < 20_000:
        raise RuntimeError("DOCX is unexpectedly small")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
    print(f"generated and validated {args.output}")


if __name__ == "__main__":
    main()
