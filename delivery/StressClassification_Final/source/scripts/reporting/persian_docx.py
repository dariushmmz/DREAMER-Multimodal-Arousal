"""Shared RTL and B Nazanin styling helpers for Persian Word deliverables."""

from __future__ import annotations

import re
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FONT_NAME = "B Nazanin"
_LTR_SPAN = re.compile(
    r"[A-Za-z0-9][A-Za-z0-9_./+×→=:%(),\[\]{}'\-–— ]*[A-Za-z0-9)\]}%]"
    r"|[A-Za-z0-9]"
)


def plain_markdown(text: str) -> str:
    """Remove the limited inline Markdown used in the report sources."""

    return text.replace("**", "").replace("`", "").replace("*", "").strip()


def set_run_font(run, size: float, *, bold: bool = False, rtl: bool = True) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)
    direction = properties.find(qn("w:rtl"))
    if direction is None:
        direction = OxmlElement("w:rtl")
        properties.append(direction)
    direction.set(qn("w:val"), "1" if rtl else "0")
    language = properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        properties.append(language)
    language.set(qn("w:val"), "fa-IR" if rtl else "en-US")
    language.set(qn("w:bidi"), "fa-IR")


def set_paragraph_rtl(
    paragraph,
    *,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_before: float = 0,
    space_after: float = 6,
    line_spacing: float = 1.25,
) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    properties = paragraph._p.get_or_add_pPr()
    bidi = properties.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        properties.append(bidi)
    bidi.set(qn("w:val"), "1")


def add_mixed_text(paragraph, text: str, size: float, *, bold: bool = False) -> None:
    """Add Persian text while explicitly marking Latin/number spans LTR."""

    cursor = 0
    for match in _LTR_SPAN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size, bold=bold, rtl=True)
        run = paragraph.add_run(match.group())
        set_run_font(run, size, bold=bold, rtl=False)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size, bold=bold, rtl=True)


def set_cell_rtl(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    direction = properties.find(qn("w:textDirection"))
    if direction is None:
        direction = OxmlElement("w:textDirection")
        properties.append(direction)
    direction.set(qn("w:val"), "lrTb")
    for paragraph in cell.paragraphs:
        set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)


def set_table_rtl(table) -> None:
    properties = table._tbl.tblPr
    bidi = properties.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        properties.append(bidi)
    bidi.set(qn("w:val"), "1")


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def add_page_number(paragraph) -> None:
    set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_mixed_text(paragraph, "صفحه ", 9)
    run = paragraph.add_run()
    set_run_font(run, 9, rtl=False)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def add_toc(paragraph) -> None:
    set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    run = paragraph.add_run()
    set_run_font(run, 11, rtl=False)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "فهرست مطالب پس از بازکردن فایل در Word به‌روزرسانی می‌شود."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])


def iter_all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def validate_b_nazanin(document, path: Path) -> None:
    missing = []
    for paragraph in iter_all_paragraphs(document):
        for run in paragraph.runs:
            if not run.text:
                continue
            properties = run._element.rPr
            fonts = None if properties is None else properties.rFonts
            values = [] if fonts is None else [fonts.get(qn(f"w:{key}")) for key in ("ascii", "hAnsi", "eastAsia", "cs")]
            if values != [FONT_NAME] * 4:
                missing.append(run.text[:30])
    if missing:
        raise RuntimeError(f"{path}: runs without B Nazanin: {missing[:5]}")
